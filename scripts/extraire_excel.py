"""Extraction en lecture seule du classeur Excel d'origine.

Produit, dans le dossier ``extraction/`` (non versionné) :

- ``vba.txt``    : le code VBA décompressé (oletools / olevba) ;
- ``feuilles.txt`` : pour chacun des 5 onglets, chaque cellule non vide avec sa
  formule et la dernière valeur calculée enregistrée par Excel (openpyxl).

Usage ::

    python scripts/extraire_excel.py [chemin_du_classeur.xlsm]

Le classeur n'est jamais modifié : openpyxl l'ouvre en lecture et rien n'est
sauvegardé.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from oletools.olevba import VBA_Parser

RACINE = Path(__file__).resolve().parents[1]
CLASSEUR_DEFAUT = RACINE / "Pricer_Options_Rayane_ALKHALAF.xlsm"


def extraire_vba(classeur: Path) -> str:
    """Renvoie le code source de tous les modules VBA du classeur."""
    parser = VBA_Parser(str(classeur))
    morceaux = []
    try:
        for _, flux, nom, code in parser.extract_macros():
            morceaux.append(f"' ===== {nom} ({flux}) =====\n{code}")
    finally:
        parser.close()
    return "\n".join(morceaux)


def extraire_feuilles(classeur: Path) -> str:
    """Renvoie, onglet par onglet, les formules et les valeurs calculées."""
    formules = openpyxl.load_workbook(classeur, keep_vba=True, read_only=True)
    valeurs = openpyxl.load_workbook(classeur, data_only=True, read_only=True)
    lignes: list[str] = []
    for ws in formules.worksheets:
        wv = valeurs[ws.title]
        lignes.append(f"===== {ws.title} =====")
        for ligne_f, ligne_v in zip(ws.iter_rows(), wv.iter_rows(), strict=False):
            for cf, cv in zip(ligne_f, ligne_v, strict=False):
                if cf.value is None:
                    continue
                texte = f"{cf.coordinate}: {cf.value!r}"
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    texte += f"  -> {cv.value!r}"
                lignes.append(texte)
    return "\n".join(lignes)


def main() -> None:
    classeur = Path(sys.argv[1]) if len(sys.argv) > 1 else CLASSEUR_DEFAUT
    sortie = RACINE / "extraction"
    sortie.mkdir(exist_ok=True)
    (sortie / "vba.txt").write_text(extraire_vba(classeur), encoding="utf-8")
    (sortie / "feuilles.txt").write_text(extraire_feuilles(classeur), encoding="utf-8")
    print(f"Extraction écrite dans {sortie}")


if __name__ == "__main__":
    main()
